#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.14"
# dependencies = ["python-pptx==1.0.2"]
# ///

"""Create a presentation from an Azure POTX/PPTX template."""

from __future__ import annotations

import argparse
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

CONTENT_TYPES_PATH = "[Content_Types].xml"
PRESENTATION_CONTENT_TYPE = (
    b"application/vnd.openxmlformats-officedocument."
    b"presentationml.presentation.main+xml"
)
TEMPLATE_CONTENT_TYPE = (
    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)
SECTION_LIST_TAG = (
    "{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst"
)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Create a PPTX from a POTX/PPTX template and add one slide. "
            "The input file is never overwritten."
        )
    )
    parser.add_argument("template", type=Path, help="Input .potx or .pptx file")
    parser.add_argument("output", type=Path, help="Output .pptx file")
    parser.add_argument(
        "--layout-index",
        type=int,
        default=12,
        help="Slide layout index to use (default: 12)",
    )
    parser.add_argument("--title", help="Text for the title placeholder")
    parser.add_argument(
        "--body",
        action="append",
        default=[],
        help="Body paragraph. Repeat the option for multiple paragraphs.",
    )
    parser.add_argument(
        "--body-placeholder-index",
        type=int,
        help="Placeholder index for body text",
    )
    slide_group = parser.add_mutually_exclusive_group()
    slide_group.add_argument(
        "--clear-slides",
        dest="clear_slides",
        action="store_true",
        default=None,
        help="Remove existing slides before adding the new slide",
    )
    slide_group.add_argument(
        "--keep-slides",
        dest="clear_slides",
        action="store_false",
        help="Keep existing slides",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Replace an existing output file",
    )
    return parser


def validate_paths(template: Path, output: Path, force: bool) -> None:
    if not template.is_file():
        raise ValueError(f"input file does not exist: {template}")
    if template.suffix.lower() not in {".potx", ".pptx"}:
        raise ValueError("input extension must be .potx or .pptx")
    if output.suffix.lower() != ".pptx":
        raise ValueError("output extension must be .pptx")
    if template.resolve() == output.resolve():
        raise ValueError("input and output must be different files")
    if output.exists() and not force:
        raise ValueError(f"output already exists; use --force to replace it: {output}")


def convert_potx_to_pptx(source: Path, destination: Path) -> None:
    converted = False
    with (
        zipfile.ZipFile(source, "r") as source_zip,
        zipfile.ZipFile(destination, "w", allowZip64=True) as destination_zip,
    ):
        if CONTENT_TYPES_PATH not in source_zip.namelist():
            raise ValueError(f"{CONTENT_TYPES_PATH} is missing from {source}")

        for entry in source_zip.infolist():
            with source_zip.open(entry) as source_file:
                content = source_file.read()

            if entry.filename == CONTENT_TYPES_PATH:
                occurrences = content.count(TEMPLATE_CONTENT_TYPE)
                if occurrences != 1:
                    raise ValueError(
                        "expected exactly one PresentationML template content type "
                        f"in {CONTENT_TYPES_PATH}, found {occurrences}"
                    )
                content = content.replace(
                    TEMPLATE_CONTENT_TYPE,
                    PRESENTATION_CONTENT_TYPE,
                    1,
                )
                converted = True

            destination_zip.writestr(entry, content)

    if not converted:
        raise ValueError("the POTX content type was not converted")


def delete_all_slides(presentation: PresentationType) -> int:
    # python-pptx has no public slide-removal API. Keep the private dependency
    # limited to _sldIdLst and pin the library version in the PEP 723 metadata.
    slide_ids = presentation.slides._sldIdLst
    removed = 0
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
        removed += 1
    return removed


def remove_section_lists(presentation: PresentationType) -> int:
    # Sections refer to slide IDs. Once every slide is removed, retaining the
    # section list leaves references to slides that no longer exist.
    presentation_element = presentation.part._element
    section_lists = list(presentation_element.iter(SECTION_LIST_TAG))
    for section_list in section_lists:
        parent = section_list.getparent()
        if parent is not None:
            parent.remove(section_list)
    return len(section_lists)


def count_section_lists(presentation: PresentationType) -> int:
    return sum(1 for _ in presentation.part._element.iter(SECTION_LIST_TAG))


def available_layouts(presentation: PresentationType) -> str:
    return ", ".join(
        f"{index}:{layout.name or '<unnamed>'}"
        for index, layout in enumerate(presentation.slide_layouts)
    )


def select_body_placeholder(
    slide: Slide,
    placeholder_index: int | None,
):
    if placeholder_index is not None:
        try:
            placeholder = slide.placeholders[placeholder_index]
        except KeyError as exc:
            raise ValueError(
                f"body placeholder index {placeholder_index} is not present"
            ) from exc
        if not placeholder.has_text_frame:
            raise ValueError(
                f"placeholder index {placeholder_index} does not accept text"
            )
        return placeholder

    title = slide.shapes.title
    title_index = (
        title.placeholder_format.idx
        if title is not None and title.is_placeholder
        else None
    )
    for placeholder in slide.placeholders:
        if (
            placeholder.has_text_frame
            and placeholder.placeholder_format.idx != title_index
        ):
            return placeholder

    raise ValueError("the selected layout has no body text placeholder")


def set_body_text(
    slide: Slide,
    paragraphs: list[str],
    placeholder_index: int | None,
) -> None:
    if not paragraphs:
        return
    if any(not paragraph.strip() for paragraph in paragraphs):
        raise ValueError("body paragraphs must not be empty")

    placeholder = select_body_placeholder(slide, placeholder_index)
    text_frame = placeholder.text_frame
    text_frame.clear()
    text_frame.paragraphs[0].text = paragraphs[0]
    for text in paragraphs[1:]:
        text_frame.add_paragraph().text = text


def create_presentation(
    input_path: Path,
    output_path: Path,
    *,
    layout_index: int,
    title: str | None,
    body: list[str],
    body_placeholder_index: int | None,
    clear_slides: bool,
) -> tuple[int, int, int]:
    presentation = Presentation(input_path)
    layout_count = len(presentation.slide_layouts)
    if layout_index < 0 or layout_index >= layout_count:
        raise ValueError(
            f"layout index {layout_index} is outside 0..{layout_count - 1}; "
            f"available layouts: {available_layouts(presentation)}"
        )

    removed = delete_all_slides(presentation) if clear_slides else 0
    removed_sections = remove_section_lists(presentation) if clear_slides else 0
    slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])

    if title is not None:
        title_shape = slide.shapes.title
        if title_shape is None:
            raise ValueError("the selected layout has no title placeholder")
        title_shape.text = title

    set_body_text(slide, body, body_placeholder_index)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)

    reopened = Presentation(output_path)
    if len(reopened.slides) == 0:
        raise ValueError("saved presentation contains no slides")
    if clear_slides and count_section_lists(reopened) != 0:
        raise ValueError("saved presentation contains stale section information")
    return removed, removed_sections, len(reopened.slides)


def run(args: argparse.Namespace) -> None:
    template = args.template.expanduser()
    output = args.output.expanduser()
    validate_paths(template, output, args.force)
    clear_slides = (
        args.clear_slides
        if args.clear_slides is not None
        else template.suffix.lower() == ".potx"
    )

    if template.suffix.lower() == ".potx":
        with tempfile.TemporaryDirectory(prefix="azure-pptx-") as temp_dir:
            converted = Path(temp_dir) / "template.pptx"
            convert_potx_to_pptx(template, converted)
            removed, removed_sections, slide_count = create_presentation(
                converted,
                output,
                layout_index=args.layout_index,
                title=args.title,
                body=args.body,
                body_placeholder_index=args.body_placeholder_index,
                clear_slides=clear_slides,
            )
    else:
        removed, removed_sections, slide_count = create_presentation(
            template,
            output,
            layout_index=args.layout_index,
            title=args.title,
            body=args.body,
            body_placeholder_index=args.body_placeholder_index,
            clear_slides=clear_slides,
        )

    print(f"created: {output}")
    print(f"removed slides: {removed}")
    print(f"removed section lists: {removed_sections}")
    print(f"output slides: {slide_count}")


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        run(args)
    except (
        KeyError,
        OSError,
        PackageNotFoundError,
        ValueError,
        zipfile.BadZipFile,
    ) as exc:
        parser.exit(1, f"error: {exc}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
